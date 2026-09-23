"""Check bilingual publication contracts and render every final page for review."""
import argparse,hashlib,io,json,re,sys,zipfile
from pathlib import Path
import numpy as np
import pdfplumber,pypdfium2 as pdfium,zxingcpp
from PIL import Image,ImageDraw
from pypdf import PdfReader
from docx import Document
ROOT=Path(__file__).resolve().parents[2];PUBLIC=ROOT/'web/public';QA=ROOT/'.production-runtime/qa-v9/publications'
def load(path):return json.loads((ROOT/path).read_text(encoding='utf-8-sig'))
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def compact(text):return re.sub(r'\s+','',text).replace('—','-').replace('Arc’teryx',"Arc'teryx")

def run(render=True):
    QA.mkdir(parents=True,exist_ok=True);manifest=load('web/src/data/downloads.json');profile=load('web/src/data/profile.json');publication=load('web/src/data/publication.json');checks=[];artifacts=[]
    def check(name,ok):checks.append({'check':name,'passed':bool(ok)})
    pdfs=[i for i in manifest['items'] if i['format']=='pdf'];check('18 bilingual PDFs',len(pdfs)==18);check('8 editable DOCX',sum(i['format']=='docx' for i in manifest['items'])==8)
    for item in pdfs:
        path=PUBLIC/item['path'].lstrip('/');reader=PdfReader(path);key=item['id'];lang=item['lang'];text='\n'.join(p.extract_text() or '' for p in reader.pages);doc=pdfium.PdfDocument(str(path));pages=[]
        check(key+' measured download',len(reader.pages)==item['pages'] and path.stat().st_size==item['bytes'] and sha(path)==item['sha256'])
        check(key+' size budget',path.stat().st_size<=20*1048576)
        check(key+' identity',profile['name'][lang].lower() in text.lower() and profile['position'][lang] in text)
        check(key+' no obsolete identity',not re.search('Design Lead|AI-native Design Technologist|设计主导',text))
        check(key+' searchable body on every page',all(len(p.extract_text() or '')>35 for p in reader.pages))
        check(key+' text encoding','\ufffd' not in text and '\x00' not in text)
        urls=[str(a.get_object().get('/A',{}).get('/URI','')) for p in reader.pages for a in p.get('/Annots',[])];site=profile['url'].rstrip('/')+('/en/' if lang=='en' else '/')
        check(key+' localized website link',site in urls)
        if item['type']=='portfolio':
            mid=('full' if item['path'].find('full-portfolio')>=0 else 'selected' if item['variant']=='overview' else item['variant'])+('' if lang=='zh' else '-en')
            source=load('deliverables/portfolio/'+mid+'-manifest.json')
            check(key+' source snapshot',all(sha(ROOT/p)==value for p,value in source['sourceHashes'].items()))
            check(key+' page budget',len(reader.pages)==source['pageCount'])
            expected=set(p.stem[:-3] for p in (ROOT/'web/src/content/works').glob('*.zh.md'))
            if mid.startswith('full'):check(key+' all 32 cases',set(source['caseSlugs'])==expected and len(source['caseSlugs'])==32)
            elif item['variant']=='overview':check(key+' homepage sequence',source['caseSlugs']==publication['homepageSelection'] and len(reader.pages)==33)
            else:check(key+' selected budget',len(reader.pages)==publication['editions'][item['variant']]['pageCount'])
            check(key+' bookmarks',len(reader.outline)>=len(source['caseSlugs']))
            check(key+' canonical case links',all('/#/work/' not in u for u in urls) and any(('/en/work/' if lang=='en' else '/work/') in u for u in urls))
        else:
            check(key+' one A4 page',len(reader.pages)==1 and abs(float(reader.pages[0].mediabox.width)-595.28)<1 and abs(float(reader.pages[0].mediabox.height)-841.89)<1)
            editable=path.with_suffix('.docx');word=Document(editable);paragraphs=list(word.paragraphs)+[p for table in word.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]
            check(key+' complete Word text',all(compact(p.text) in compact(text) for p in paragraphs if p.text.strip()))
            check(key+' shared experience',all(compact(b) in compact(text) for bullets in profile['resume']['experienceLocalized'][lang].values() for b in bullets))
            check(key+' original portrait', (PUBLIC/'media/v7/resume-portrait.png').read_bytes() in [zipfile.ZipFile(editable).read(n) for n in zipfile.ZipFile(editable).namelist() if n.startswith('word/media/')])
        # Physical bounds are measured from PDF glyphs, independent of the generator.
        with pdfplumber.open(path) as pdoc:
            outside=[];small_body=[]
            for index,page in enumerate(pdoc.pages):
                chars=[c for c in page.chars if c['text'].strip()]
                if any(c['x0']<16 or c['x1']>page.width-16 or c['top']<12 or c['bottom']>page.height-12 for c in chars):outside.append(index+1)
                if item['type']=='resume':check(key+' body at least 10.5pt',sum(c['size']>=10.49 for c in chars)>len(chars)*.7)
            check(key+' glyphs inside page',not outside)
        for i in range(len(doc)):
            page=doc[i]
            if render:
                image=page.render(scale=1.5 if item['type']=='resume' else 1.25).to_pil().convert('RGB')
                dest=QA/path.stem;dest.mkdir(exist_ok=True);target=dest/f'{i+1:03d}.png';image.save(target);pages.append(target)
                if i==0:
                    preview=PUBLIC/item['preview'].lstrip('/');preview.parent.mkdir(parents=True,exist_ok=True)
                    image.thumbnail((640,1000),Image.Resampling.LANCZOS);image.save(preview,'WEBP',quality=88)
            if i==len(doc)-1:
                qr_image=page.render(scale=3).to_pil().convert('RGB');codes=[x.text for x in zxingcpp.read_barcodes(np.asarray(qr_image))];check(key+' rendered QR',site in codes)
            page.close()
        doc.close()
        if pages and item['type']=='portfolio':
            for start in range(0,len(pages),20):
                sheet=Image.new('RGB',(1600,1350),'#dedede');draw=ImageDraw.Draw(sheet)
                for j,file in enumerate(pages[start:start+20]):
                    with Image.open(file) as im:
                        im.thumbnail((380,238),Image.Resampling.LANCZOS);x=(j%4)*400+10;y=(j//4)*270+20
                        sheet.paste(im,(x,y));draw.text((x,y-15),f'{start+j+1:03d}',fill='black')
                sheet.save(QA/(path.stem+f'-sheet-{start//20+1:02d}.jpg'),quality=88)
        artifacts.append({'path':item['path'],'pages':len(reader.pages),'lang':lang,'renderedPages':len(pages)})
        print(f'{key}: {len(reader.pages)} pages inspected'+(' and rendered' if render else ''),flush=True)
    check('matching public manifest',sha(ROOT/'web/src/data/downloads.json')==sha(PUBLIC/'downloads/manifest.json'))
    result={'passed':all(c['passed'] for c in checks),'checks':checks,'artifacts':artifacts,'visualReview':'Rendered page images require human visual inspection; structural checks do not substitute for it.'}
    (QA/'report.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    failures=[c['check'] for c in checks if not c['passed']];print(json.dumps({'passed':result['passed'],'checks':len(checks),'failures':failures},ensure_ascii=False))
    if failures:raise SystemExit(1)
if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8');p=argparse.ArgumentParser();p.add_argument('--no-render',action='store_true');args=p.parse_args();run(not args.no_render)
