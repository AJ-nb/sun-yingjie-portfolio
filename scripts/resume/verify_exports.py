"""Verify final PDFs/DOCX, source snapshots and decoded QR links.

Requires pypdf, python-docx, pypdfium2, numpy and zxing-cpp. Visual review remains
a separate required step: this report does not claim image or layout approval.
"""
import hashlib,json,re,sys,zipfile,io
from pathlib import Path
from docx import Document
from pypdf import PdfReader
from PIL import Image
import numpy as np
import pypdfium2 as pdfium
try:
    import zxingcpp
except ImportError as error:
    raise SystemExit('QR validation requires zxing-cpp in the selected Python runtime. Install it before verifying; do not substitute source-string checks.') from error

ROOT=Path(__file__).resolve().parents[2]
def digest(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def compact(value):return re.sub(r'\s+','',value).replace('Arc’teryx',"Arc'teryx")
def qr_from_pdf(path,page_number):
    document=pdfium.PdfDocument(str(path));page=document[page_number]
    image=page.render(scale=3).to_pil().convert('RGB')
    results=zxingcpp.read_barcodes(np.asarray(image))
    page.close();document.close()
    return [result.text for result in results]
def urls(reader):
    result=[]
    for page in reader.pages:
        for ref in page.get('/Annots',[]):
            action=ref.get_object().get('/A',{})
            if action.get('/URI'):result.append(str(action['/URI']))
    return result

def run():
    profile=json.loads((ROOT/'web/src/data/profile.json').read_text(encoding='utf-8'))
    publication=json.loads((ROOT/'web/src/data/publication.json').read_text(encoding='utf-8'))
    mapping=json.loads((ROOT/'web/src/data/publication-pages.json').read_text(encoding='utf-8'))
    checks=[];artifacts=[]
    def check(label,condition):checks.append({'check':label,'passed':bool(condition)})
    def snapshot(label,manifest,allowed_stale=()):
        stale=[p for p,value in manifest['sourceHashes'].items() if not (ROOT/p).is_file() or digest(ROOT/p)!=value]
        check(label+' source snapshot matches',not stale or set(stale).issubset(set(allowed_stale)))
    public_cases={path.name[:-6] for path in (ROOT/'web/src/content/works').glob('*.zh.md')} - set(publication['excludedSlugs'])
    deferred_portfolio_cases=set(publication.get('portfolioPdfDeferredSlugs',[]))
    snapshot_exemptions=publication.get('portfolioPdfSnapshotExemptions',{})
    for variant,manifest_id in [('overview','selected'),('brand','brand'),('physical','physical'),('digital','digital'),('full','full')]:
        manifest=json.loads((ROOT/f'deliverables/portfolio/{manifest_id}-manifest.json').read_text(encoding='utf-8'))
        path=ROOT/'web/public'/manifest['publicPath'].lstrip('/');reader=PdfReader(path)
        texts=[p.extract_text() or '' for p in reader.pages]
        check(variant+' page count',len(reader.pages)==manifest['pageCount'])
        check(variant+' case count',len(manifest['caseSlugs'])==manifest['caseCount'])
        if variant=='full':check('full includes every non-deferred public case',set(manifest['caseSlugs'])==public_cases-deferred_portfolio_cases)
        else:
            edition=publication['editions'][variant]
            check(variant+' planned page count',len(reader.pages)==edition['pageCount']<=edition['maxPages'])
            check(variant+' ordered cases',manifest['caseSlugs']==[x['slug'] for x in edition['cases']])
        if variant=='brand':check('brand contains only four authorized cases',set(manifest['caseSlugs'])=={'hermes','arcteryx','periastra','yelisi'})
        check(variant+' public hash',digest(path)==manifest['publicSha256'])
        check(variant+' below 20 MiB',path.stat().st_size<20*1024*1024)
        check(variant+' searchable Chinese',profile['name']['zh'] in texts[0])
        check(variant+' text on every page',all(len(t.strip())>20 for t in texts))
        check(variant+' no replacement characters',all('\ufffd' not in t for t in texts))
        check(variant+' no leaked markup',all(not any(mark in t for mark in ['<details>','<summary>','###']) for t in texts))
        snapshot(variant,manifest,snapshot_exemptions.get(manifest_id,[]))
        check(variant+' rendering font snapshot',all(Path(p).is_file() and digest(Path(p))==value for p,value in manifest['externalFontHashes'].items()))
        check(variant+' original media hashes',all(digest(ROOT/'web/public'/m['publicPath'].lstrip('/'))==m['sha256'] for m in manifest['media']))
        expected_url=profile['url'];link_list=urls(reader)
        check(variant+' current canonical link',expected_url in link_list and manifest['url']==expected_url)
        check(variant+' all case links',all(expected_url.rstrip('/')+'/#/work/'+slug in link_list for slug in manifest['caseSlugs']))
        check(variant+' final-page QR decodes canonical URL',expected_url in qr_from_pdf(path,len(reader.pages)-1))
        for record in manifest['pages']:
            if record['kind']!='evidence':continue
            page_text=compact(texts[record['page']-1])
            check(f'{variant} p{record["page"]} mapped captions',all(compact(image['caption']) in page_text for image in record['images']))
            if variant!='full':
                spec=next(p for p in mapping['casePages'][record['case']] if p['id']==record['pageId'])
                check(f'{variant} p{record["page"]} image/body mapping',record['images']==spec['images'] and record['body']==spec['body'])
        compression=json.loads((ROOT/f'deliverables/portfolio/{manifest_id}-compression.json').read_text(encoding='utf-8'))
        check(variant+' compressed text and links preserved',all(compression['checks'].values()))
        artifacts.append({'path':manifest['publicPath'],'pages':len(reader.pages),'bytes':path.stat().st_size,'sha256':digest(path)})

    resumes=json.loads((ROOT/'deliverables/resume/resume-manifest.json').read_text(encoding='utf-8'))
    check('four resume variants',set(e['id'] for e in resumes['variants'])=={'overview','brand','physical','digital'})
    check('resume body is at least 10 pt',resumes['bodyFontPt']>=10)
    for entry in resumes['variants']:
        label='resume '+entry['id'];docx=ROOT/entry['docx'];pdf=ROOT/entry['pdf'];doc=Document(docx);reader=PdfReader(pdf)
        paragraphs=list(doc.paragraphs)+[p for t in doc.tables for row in t.rows for cell in row.cells for p in cell.paragraphs]
        doc_text='\n'.join(p.text for p in paragraphs);pdf_text='\n'.join(p.extract_text() or '' for p in reader.pages)
        check(label+' one page',len(reader.pages)==1)
        check(label+' A4',abs(float(reader.pages[0].mediabox.width)-595.28)<1 and abs(float(reader.pages[0].mediabox.height)-841.89)<1)
        check(label+' semantic title',any(p.style.name=='Title' and profile['name']['zh'] in p.text for p in paragraphs))
        check(label+' editable paragraphs',len(doc.paragraphs)>=18)
        check(label+' photo and QR',len(doc.inline_shapes)==2)
        for key in ['phone','email']:check(label+' shared '+key,profile[key] in doc_text and profile[key] in pdf_text)
        for item in profile['timeline']:
            if item['id']=='ai':continue
            check(label+' period '+item['id'],compact(item['period']['zh'].replace('—','-')) in compact(doc_text))
        for key,bullets in profile['resume']['experience'].items():
            check(label+' complete experience '+key,all(compact(bullet) in compact(doc_text) and compact(bullet) in compact(pdf_text) for bullet in bullets))
        for index,p in enumerate(doc.paragraphs):
            if p.text and p.style.name=='Normal':check(label+' readable paragraph '+str(index),all((r.font.size.pt if r.font.size else 10)>=10 for r in p.runs if r.text and not r.text.startswith('\t')))
        check(label+' canonical hyperlink',profile['url'].rstrip('/') in [url.rstrip('/') for url in urls(reader)])
        check(label+' final PDF QR decodes',profile['url'] in qr_from_pdf(pdf,0))
        with zipfile.ZipFile(docx) as archive:
            values=[]
            for name in archive.namelist():
                if name.startswith('word/media/'):
                    with Image.open(io.BytesIO(archive.read(name))) as image:values += [code.text for code in zxingcpp.read_barcodes(np.asarray(image.convert('RGB')))]
            check(label+' embedded DOCX QR decodes',profile['url'] in values)
        snapshot(label,entry)
        for extension in ['pdf','docx']:
            source=pdf if extension=='pdf' else docx;public=ROOT/'web/public/downloads'/source.name
            check(label+' public '+extension+' identical',digest(source)==digest(public))
            artifacts.append({'path':'/downloads/'+source.name,'pages':1,'bytes':source.stat().st_size,'sha256':digest(source)})

    download=json.loads((ROOT/'web/src/data/downloads.json').read_text(encoding='utf-8'))
    check('download manifest contains thirteen current artifacts',download['edition']=='v7' and len(download['items'])==13)
    for item in download['items']:
        target=ROOT/'web/public'/item['path'].lstrip('/')
        check(item['id']+' measured metadata',digest(target)==item['sha256'] and target.stat().st_size==item['bytes'])
        if item['format']=='pdf':check(item['id']+' measured page count',len(PdfReader(target).pages)==item['pages'])
    result={'passed':all(c['passed'] for c in checks),'checkCount':len(checks),'checks':checks,'artifacts':artifacts,'visualReview':'Required separately. Automated extraction, QR decoding and source checks do not substitute for page-image inspection.'}
    target=ROOT/'deliverables/document-verification.json';target.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'passed':result['passed'],'checks':len(checks),'failures':[c['check'] for c in checks if not c['passed']]},ensure_ascii=False))
    if not result['passed']:raise SystemExit(1)
if __name__=='__main__':
    if hasattr(sys.stdout,'reconfigure'):sys.stdout.reconfigure(encoding='utf-8')
    run()
