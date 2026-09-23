"""Validate resume content, links, editability and actual rendered page bounds."""
import hashlib
import io
import json
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pdfplumber
import zxingcpp
from docx import Document
from PIL import Image
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[2]

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def compact(text):
    return re.sub(r'\s+', '', text).replace('Arc’teryx', "Arc'teryx").replace('—', '-')

def run():
    manifest=json.loads((ROOT/'deliverables/resume/resume-manifest.json').read_text(encoding='utf-8'))
    previews=json.loads((ROOT/'deliverables/resume/resume-previews-v77.json').read_text(encoding='utf-8'))
    profile=json.loads((ROOT/'web/src/data/profile.json').read_text(encoding='utf-8-sig'))
    positioning=json.loads((ROOT/'web/src/data/resume-positioning.json').read_text(encoding='utf-8-sig'))
    downloads=json.loads((ROOT/'web/src/data/downloads.json').read_text(encoding='utf-8-sig'))
    checks=[]
    def check(label, condition):
        checks.append({'check':label,'passed':bool(condition)})
    check('four supported role variants', {entry['id'] for entry in manifest['variants']}=={'overview','brand','physical','digital'})
    check('body type at least 10 pt', manifest['bodyFontPt']>=10)
    for entry in manifest['variants']:
        key=entry['id']; docx=ROOT/entry['docx']; pdf=ROOT/entry['pdf']; doc=Document(docx); reader=PdfReader(pdf)
        paragraphs=list(doc.paragraphs)+[p for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]
        doc_text='\n'.join(p.text for p in paragraphs)
        pdf_text='\n'.join(page.extract_text() or '' for page in reader.pages)
        direction=positioning['variants'][key]
        check(key+' no standalone project section',all(label not in doc_text and label not in pdf_text for label in ['项目实践','代表项目','设计与交付能力']))
        headings=[p.text for p in doc.paragraphs if p.style.name=='Heading 1']
        check(key+' ordered professional structure',headings==['核心优势','工作经历','专业工具','教育背景'])
        check(key+' three role-specific strengths',len(direction['strengths'])==3 and all(compact(item['label']+'：'+item['body']) in compact(doc_text) and compact(item['body']) in compact(pdf_text) for item in direction['strengths']))
        check(key+' explicit evidence for each strength',all(item['evidence'] and all(any(record['id']==evidence for record in profile['timeline']) or (ROOT/f'web/src/content/works/{evidence}.zh.md').is_file() for evidence in item['evidence']) for item in direction['strengths']))
        check(key+' role-specific summary',compact(direction['summary']) in compact(doc_text) and compact(direction['summary']) in compact(pdf_text))
        check(key+' complete professional tool groups',all(compact(item['label']+'：'+item['body']) in compact(doc_text) and compact(item['body']) in compact(pdf_text) for item in direction['tools']))
        check(key+' four targeted tool groups',len(direction['tools'])==4)
        check(key+' new confirmed tools present',all(tool in doc_text and tool in pdf_text for tool in ['Photoshop','Illustrator','Codex','ComfyUI']))
        check(key+' AIGC and open-source workflow',compact(direction['workflow']) in compact(doc_text) and compact(direction['workflow']) in compact(pdf_text) and '开源' in direction['workflow'])
        check(key+' user-confirmed tool provenance',set(positioning['confirmedInputs']['tools'])=={'Adobe Photoshop','Adobe Illustrator','Codex','ComfyUI'})
        check(key+' requested Liling period',('2026.02' in compact(doc_text) and '至今' in compact(doc_text) and '2026.02' in compact(pdf_text) and ('至今' in compact(pdf_text) or 'Present' in compact(pdf_text))))
        check(key+' independent AI practice boundary',all(token in doc_text and token in pdf_text for token in ['Lensflow', 'Resume Formatter']))
        check(key+' no invented tool proficiency',all(word not in doc_text for word in ['精通','认证','熟练度']))
        check(key+' one A4 page',len(reader.pages)==1 and abs(float(reader.pages[0].mediabox.width)-595.28)<1 and abs(float(reader.pages[0].mediabox.height)-841.89)<1)
        check(key+' native Word PDF', 'Microsoft' in str(reader.metadata.producer) or 'Word' in str(reader.metadata.creator))
        check(key+' unchanged employment evidence', all(compact(bullet) in compact(doc_text) and compact(bullet) in compact(pdf_text) for bullets in profile['resume']['experience'].values() for bullet in bullets))
        check(key+' preserved identity and contacts',all(profile[field] in doc_text and profile[field] in pdf_text for field in ['phone','email']) and profile['name']['zh'] in pdf_text)
        check(key+' preserved employment and education dates',all(compact(item['period']['zh']) in compact(doc_text) and compact(item['period']['zh']) in compact(pdf_text) for item in profile['timeline'] if item['id']!='ai'))
        check(key+' no missing or replacement text', '\ufffd' not in pdf_text and all(compact(p.text) in compact(pdf_text) for p in paragraphs if p.text.strip()))
        check(key+' editable body and semantic name',len(doc.paragraphs)>=18 and any(p.style.name=='Title' and profile['name']['zh'] in p.text for p in paragraphs))
        check(key+' no italic role subtitle',doc.styles['Subtitle'].font.italic is False)
        check(key+' actual normal body at least 10 pt',all((run.font.size.pt if run.font.size else manifest['bodyFontPt'])>=10 for p in doc.paragraphs if p.text.strip() and p.style.name=='Normal' for run in p.runs if run.text.strip() and not run.text.startswith('\t')))
        check(key+' photo and QR descriptions',len(doc.inline_shapes)==2 and all(shape._inline.docPr.get('descr') for shape in doc.inline_shapes))
        links=[str(annotation.get_object().get('/A',{}).get('/URI','')) for page in reader.pages for annotation in page.get('/Annots',[])]
        check(key+' active portfolio link', profile['url'].rstrip('/') in [link.rstrip('/') for link in links])
        embedded=[]
        with zipfile.ZipFile(docx) as archive:
            media=[archive.read(name) for name in archive.namelist() if name.startswith('word/media/')]
            portrait=(ROOT/'web/public/media/v7/resume-portrait.png').read_bytes()
            check(key+' original portrait preserved',portrait in media)
            for data in media:
                with Image.open(io.BytesIO(data)) as im:
                    embedded.extend(code.text for code in zxingcpp.read_barcodes(np.asarray(im.convert('RGB'))))
        check(key+' DOCX QR decodes',profile['url'] in embedded)
        preview=next(record for record in previews['pages'] if record['id']==key)
        with Image.open(ROOT/preview['qaPage']) as im:
            qr_values=[code.text for code in zxingcpp.read_barcodes(np.asarray(im.convert('RGB')))]
        check(key+' rendered QR decodes',profile['url'] in qr_values)
        check(key+' preview generated from final PDF',preview['pdfSha256']==digest(pdf) and preview['previewSha256']==digest(ROOT/preview['preview']))
        with Image.open(ROOT/preview['preview']) as im:
            check(key+' full A4 preview',im.width==640 and abs(im.height/im.width-297/210)<0.005)
        with pdfplumber.open(pdf) as pages:
            chars=pages.pages[0].chars
            # Word permits hanging CJK punctuation beyond the right text margin;
            # require a physical safety margin instead of rejecting that typesetting.
            # Word emits a trailing whitespace glyph after the final table on
            # some one-page exports. It is invisible layout padding rather
            # than resume content, so only visible characters participate in
            # the physical text-bound check.
            visible_chars=[char for char in chars if char['text'].strip()]
            check(key+' all text inside safe page bounds',all(char['x0']>=40 and char['x1']<=560 and char['top']>=35 and char['bottom']<=800 for char in visible_chars))
            # Native Word PDF font sizes are rounded to 0.12 pt increments.
            check(key+' actual PDF body font size',sum(abs(char['size']-manifest['bodyFontPt'])<0.1 for char in chars)>300)
        check(key+' current source snapshot',all((ROOT/path).is_file() and digest(ROOT/path)==sha for path,sha in entry['sourceHashes'].items()))
        check(key+' manifest artifact hashes',digest(pdf)==entry['pdfSha256'] and digest(docx)==entry['docxSha256'])
        for extension in ['pdf','docx']:
            source=pdf if extension=='pdf' else docx
            check(key+' public '+extension+' matches',digest(source)==digest(ROOT/'web/public/downloads'/source.name))
            metadata=next(item for item in downloads['items'] if item['id']=='resume-'+key+'-'+extension)
            check(key+' '+extension+' download metadata matches',metadata['sha256']==digest(source) and metadata['bytes']==source.stat().st_size and metadata['pages']==1)
    check('public download manifest matches source',digest(ROOT/'web/src/data/downloads.json')==digest(ROOT/'web/public/downloads/manifest.json'))
    result={'edition':'resume-v7.7','checkedAt':datetime.now(timezone.utc).isoformat(),'passed':all(item['passed'] for item in checks),'checkCount':len(checks),'checks':checks,'renderer':{'actual':'Microsoft Word native fixed-format PDF export, then bundled Poppler 150 dpi','skillRenderer':'render_docx.py was attempted with runtime-only PATH and failed because this Windows bundle has no LibreOffice. No system LibreOffice was used.','log':'deliverables/resume/qa/v77/skill-render.log'},'visualReview':'Separate manual review of all four final page images is required.'}
    (ROOT/'deliverables/resume/resume-verification-v77.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    failures=[item['check'] for item in checks if not item['passed']]
    print(json.dumps({'passed':result['passed'],'checks':len(checks),'failures':failures},ensure_ascii=False))
    if failures:
        raise SystemExit(1)

if __name__=='__main__':
    run()
