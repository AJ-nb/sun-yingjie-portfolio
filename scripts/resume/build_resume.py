"""Create four editable A4 resumes with role-specific, evidence-backed strengths."""
import hashlib
import json
import sys
import re
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.graphics.barcode import qrencoder
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'deliverables/resume'
VARIANT_IDS = {'general': 'overview', 'brand': 'brand', 'product': 'physical', 'digital': 'digital'}
BODY_PT = 11
BODY_LEADING_PT = 17
INK = '20221F'
MUTED = '656962'

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def make_qr(value, output):
    qr = qrencoder.QRCode(1, qrencoder.QRErrorCorrectLevel.M)
    qr.addData(value); qr.version = qr.calculate_version(); qr.make()
    modules = qr.getModuleCount(); scale = 12
    im = Image.new('RGB', ((modules + 8) * scale, (modules + 8) * scale), 'white')
    draw = ImageDraw.Draw(im)
    for y in range(modules):
        for x in range(modules):
            if qr.isDark(y, x):
                draw.rectangle(((x+4)*scale, (y+4)*scale, (x+5)*scale-1, (y+5)*scale-1), fill='#20221f')
    im.save(output)

def build_one(profile, variant, positioning, qr_path, lang="zh"):
    zh = lang == "zh"
    t = lambda cn,en: cn if zh else en
    variant_id = VARIANT_IDS[variant['id']]
    suffix = '' if variant_id == 'overview' else '-' + variant_id
    filename = 'sun-yingjie-resume' + suffix + ('' if zh else '-en')
    timeline = {item['id']: item for item in profile['timeline']}
    resume = profile['resume']
    direction = positioning['variants'][variant_id]['localized'][lang]
    doc = Document(); section = doc.sections[0]
    section.page_width = Mm(210); section.page_height = Mm(297)
    # Keep the one-page A4 contract while accommodating the expanded evidence
    # language in the role variants. The visible content remains within the
    # existing safe bounds verified by verify_resumes.py.
    section.top_margin = Mm(15); section.bottom_margin = Mm(10)
    section.left_margin = Mm(17); section.right_margin = Mm(17)
    section.header_distance = Mm(4); section.footer_distance = Mm(4)
    for style_name in ['Normal', 'Title', 'Heading 1', 'Subtitle']:
        style = doc.styles[style_name]
        style.font.name = 'Microsoft YaHei' if zh else 'Aptos'; style.font.size = Pt(BODY_PT)
        style.font.italic = False
        style.font.color.rgb = RGBColor.from_string('000000')
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = Pt(BODY_LEADING_PT)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for grid in list(section._sectPr.findall(qn('w:docGrid'))):
        section._sectPr.remove(grid)

    def para(text='', size=BODY_PT, bold=False, before=0, after=2, style=None, parent=doc, color=INK):
        p = parent.add_paragraph(style=style)
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = Pt(BODY_LEADING_PT if size <= 11 else size * 1.25)
        p.paragraph_format.widow_control = True
        snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0')
        p._p.get_or_add_pPr().append(snap)
        r = p.add_run(text.replace('—', '-').replace('Arc’teryx', "Arc'teryx"))
        r.font.size = Pt(size); r.bold = bold
        r.font.color.rgb = RGBColor.from_string(color)
        return p

    def heading(text):
        p = para(text, 11.5, True, 10, 4, 'Heading 1', color='000000')
        p.paragraph_format.keep_with_next = True

    def dated_title(text, date):
        p = para(text, 10.5, True, 3, 2)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(176), WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run('\t' + date.replace('—', '-')); r.font.size = Pt(9.5); r.bold = False
        r.font.color.rgb = RGBColor.from_string(MUTED)
        p.paragraph_format.keep_with_next = True

    def label_body(text, after=2):
        label, separator, body = text.partition('：')
        p = para('', after=after)
        if separator:
            r = p.add_run(label + t('：', ': ')); r.bold = True; r.font.size = Pt(BODY_PT)
            r.font.color.rgb = RGBColor.from_string(INK)
        else:
            body = text
        r = p.add_run(body); r.font.size = Pt(BODY_PT); r.font.color.rgb = RGBColor.from_string(INK)
        return p

    def clear_cell_chrome(table):
        # Retain the inherited portrait/QR alignment, with no visible grid.
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell in table.rows[0].cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                node = OxmlElement('w:' + side); node.set(qn('w:w'), '0'); node.set(qn('w:type'), 'dxa'); margins.append(node)
            tc_pr.append(margins)

    # This is an inherited photo header, not a data table.
    header = doc.add_table(rows=1, cols=2); header.autofit = False
    header.columns[0].width = Mm(146); header.columns[1].width = Mm(30)
    left, right = header.rows[0].cells
    left.width = Mm(146); right.width = Mm(30)
    clear_cell_chrome(header)
    for cell in (left, right):
        for p in list(cell.paragraphs): p._element.getparent().remove(p._element)
    para(profile['name'][lang], 29, False, 0, 4, 'Title', left, color='000000')
    para(variant['title'][lang], 11.5, False, 0, 6, 'Subtitle', left, color='000000')
    para(profile['phone'] + '    ' + profile['email'], 9.5, after=0, parent=left, color=MUTED)
    # Resume portrait keeps the canonical 1:1 source background; website and
    # portfolio PDFs use the deterministic transparent derivative.
    portrait = ROOT/'web/public/media/v7/resume-portrait.png'
    if not portrait.is_file(): raise FileNotFoundError(portrait)
    p = right.add_paragraph(); p.paragraph_format.line_spacing = 1; p.paragraph_format.space_after=Pt(0); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    portrait_shape = p.add_run().add_picture(str(portrait), width=Mm(27), height=Mm(27))
    portrait_shape._inline.docPr.set('descr', profile['name']['zh'] + '个人照片')
    para(direction['summary'], before=10, after=0)

    heading(t('核心优势','Design focus'))
    for strength in direction['strengths']:
        label_body(strength['label'] + '：' + strength['body'], after=3)

    heading(t('工作经历','Experience'))
    for key in ['liling', 'hannstar', 'benwu', 'ouyin']:
        item = timeline[key]
        dated_title(item['place'][lang] + '  |  ' + item['role'][lang], re.sub(r'（.*?）| \(record.*?\)', '', item['period'][lang]))
        # Shared experience bullets are identical across variants.
        for bullet in resume['experienceLocalized'][lang][key]: para(bullet, after=2)

    heading(t('专业工具','Tools and independent practice'))
    for group in direction['tools']:
        label_body(group['label'] + '：' + group['body'], after=2)
    para(direction['workflow'], before=4, after=0)
    boundary = direction.get('projectBoundary')
    if boundary:
        label_body(('AI / 开源实践：' if variant_id != 'digital' else 'AI / open-source practice: ') + boundary, after=1)

    heading(t('教育背景','Education'))
    education = timeline['education']
    dated_title(education['place'][lang] + '  |  ' + t('产品设计本科','BA Product Design'), education['period'][lang])

    footer = doc.add_table(rows=1, cols=2); footer.autofit=False
    footer.columns[0].width=Mm(22); footer.columns[1].width=Mm(154)
    footer.cell(0,0).width=Mm(22); footer.cell(0,1).width=Mm(154)
    clear_cell_chrome(footer)
    p=footer.cell(0,0).paragraphs[0];p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(0)
    qr_shape=p.add_run().add_picture(str(qr_path), width=Mm(19));qr_shape._inline.docPr.set('descr', '扫码查看孙英杰完整作品集 '+(profile['url'] + ('/en/' if not zh else '/')))
    p=footer.cell(0,1).paragraphs[0];p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(2)
    r=p.add_run(t('作品集  /  完整案例与设计过程','Portfolio / complete cases and design process'));r.font.size=Pt(10);r.font.color.rgb=RGBColor.from_string(INK)
    p=footer.cell(0,1).add_paragraph();p.paragraph_format.space_after=Pt(0)
    hyperlink=OxmlElement('w:hyperlink');hyperlink.set(qn('r:id'),p.part.relate_to((profile['url'] + ('/en/' if not zh else '/')),RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');properties=OxmlElement('w:rPr');size=OxmlElement('w:sz');size.set(qn('w:val'),'19');properties.append(size);color=OxmlElement('w:color');color.set(qn('w:val'), MUTED);properties.append(color);run.append(properties)
    text=OxmlElement('w:t');text.text=(profile['url'] + ('/en/' if not zh else '/'));run.append(text);hyperlink.append(run);p._p.append(hyperlink)
    doc.core_properties.author=profile['name'][lang];doc.core_properties.title=profile['name'][lang]+' '+variant['title'][lang]
    doc.core_properties.subject=t('工业与产品设计 职业档案 2026.09','Industrial and product design professional record 2026.09')
    output=OUT/(filename+'.docx');doc.save(output)
    content={'variant':variant_id,'title':variant['title'],'summary':direction['summary'],'strengths':direction['strengths'],'experience':resume['experienceLocalized'][lang],'tools':direction['tools'],'workflow':direction['workflow'],'workflowEvidence':direction['workflowEvidence'],'confirmedInputs':positioning['confirmedInputs'],'sections':['核心优势','工作经历','专业工具','教育背景'],'url':(profile['url'] + ('/en/' if not zh else '/'))}
    content_path=OUT/(filename+'-content.json');content_path.write_text(json.dumps(content,ensure_ascii=False,indent=2),encoding='utf-8')
    if variant_id == 'overview':
        (OUT/'resume-content.json').write_text(json.dumps(content,ensure_ascii=False,indent=2),encoding='utf-8')
    evidence_ids={evidence for strength in direction['strengths'] for evidence in strength['evidence']}
    evidence_ids.update(direction['workflowEvidence'])
    case_sources=[ROOT/f'web/src/content/works/{slug}.{lang}.md' for slug in sorted(evidence_ids) if slug not in timeline and not slug.startswith('user-confirmed-')]
    sources=[ROOT/'web/src/data/profile.json',ROOT/'web/src/data/resume-positioning.json',ROOT/'scripts/resume/build_resume.py',portrait,qr_path,*case_sources]
    return {'id':variant_id,'lang':lang,'stem':filename,'docx':output.relative_to(ROOT).as_posix(),'content':content_path.relative_to(ROOT).as_posix(),'sourceHashes':{p.relative_to(ROOT).as_posix():digest(p) for p in sources},'docxSha256':digest(output)}

def build():
    OUT.mkdir(parents=True,exist_ok=True)
    profile=json.loads((ROOT/'web/src/data/profile.json').read_text(encoding='utf-8-sig'))
    positioning=json.loads((ROOT/'web/src/data/resume-positioning.json').read_text(encoding='utf-8-sig'))
    entries=[]
    for lang in ['zh','en']:
        qr_path=OUT/('portfolio-qr'+('' if lang=='zh' else '-en')+'.png')
        make_qr(profile['url']+('/en/' if lang=='en' else '/'),qr_path)
        for variant in profile['resume']['variants']:
            entries.append(build_one(profile,variant,positioning,qr_path,lang))
    (OUT/'resume-build.json').write_text(json.dumps({'edition':positioning['edition'],'url':profile['url'],'bodyFontPt':BODY_PT,'bodyLeadingPt':BODY_LEADING_PT,'variants':entries},ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'variants':len(entries),'output':str(OUT)},ensure_ascii=False))

if __name__=='__main__':
    if hasattr(sys.stdout,'reconfigure'):sys.stdout.reconfigure(encoding='utf-8')
    build()
